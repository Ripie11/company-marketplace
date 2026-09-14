#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
extract_docx.py - 从 .docx 文件提取格式与文本

用法:
    python extract_docx.py --input <文件或目录> --output <JSON输出文件>

依赖:
    pip install python-docx
"""
import argparse
import json
import os
import sys
from pathlib import Path
from docx import Document
from docx.shared import Pt, Mm
from docx.oxml.ns import qn


def emu_to_pt(emu):
    """EMU 转 Pt (1 pt = 12700 EMU)"""
    return round(emu / 12700, 2) if emu else None


def emu_to_mm(emu):
    """EMU 转 mm (1 mm = 36000 EMU)"""
    return round(emu / 36000, 2) if emu else None


def eastasia(run):
    """取 run 的中文字体名（w:eastAsia）"""
    rPr = run._element.rPr
    if rPr is None or rPr.rFonts is None:
        return None
    return rPr.rFonts.get(qn("w:eastAsia"))


def extract_format(doc):
    """提取文档格式

    注意：不要只看"第一个有内容的段落"——领导讲话稿首段常是主标题、
    称呼或标题混排，会把 body 判错。此处改为对前若干段做**众数统计**，
    结果更稳定（该口径来自对集团真实样张的实测校正）。
    """
    from collections import Counter

    fmt = {"fonts": {}, "paragraph": {}, "page": {}, "headings": []}

    font_counter = Counter()
    size_counter = Counter()
    indent_counter = Counter()
    spacing_counter = Counter()
    align_counter = Counter()

    scanned = 0
    for para in doc.paragraphs:
        if not para.text.strip():
            continue
        scanned += 1
        if scanned > 60:
            break
        for r in para.runs:
            if not r.text.strip():
                continue
            key = eastasia(r) or r.font.name
            if key:
                font_counter[key] += 1
            if r.font.size:
                size_counter[r.font.size.pt] += 1
        pf = para.paragraph_format
        if pf.first_line_indent:
            indent_counter[round(pf.first_line_indent.cm, 2)] += 1
        if pf.line_spacing and pf.line_spacing.pt is not None:
            spacing_counter[(round(pf.line_spacing.pt, 2), str(pf.line_spacing_rule))] += 1
        if para.alignment is not None:
            align_counter[str(para.alignment)] += 1

    if font_counter:
        fmt["fonts"]["body"] = {
            "family": font_counter.most_common(1)[0][0],
            "size_pt": size_counter.most_common(1)[0][0] if size_counter else 16,
            "color": "000000",
        }
    if spacing_counter:
        ls, ls_rule = spacing_counter.most_common(1)[0][0]
        fmt["paragraph"] = {
            "alignment": align_counter.most_common(1)[0][0] if align_counter else "justify",
            "line_spacing": ls,
            "line_spacing_rule": ls_rule,
            "first_line_indent_cm": (
                indent_counter.most_common(1)[0][0] if indent_counter else None
            ),
            "space_before_pt": 0,
            "space_after_pt": 0,
        }

    # 分布明细，便于人工核对"是否只有一种正文版式"
    fmt["fonts_dist"] = font_counter.most_common(6)
    fmt["sizes_dist"] = size_counter.most_common(6)
    fmt["indent_dist_cm"] = indent_counter.most_common(5)

    # 页面设置
    for section in doc.sections:
        fmt["page"] = {
            "paper_width_mm": emu_to_mm(section.page_width),
            "paper_height_mm": emu_to_mm(section.page_height),
            "margin_top_mm": emu_to_mm(section.top_margin),
            "margin_bottom_mm": emu_to_mm(section.bottom_margin),
            "margin_left_mm": emu_to_mm(section.left_margin),
            "margin_right_mm": emu_to_mm(section.right_margin),
        }
        break

    # 标题样式
    for para in doc.paragraphs:
        if para.style.name.startswith("Heading"):
            level = para.style.name
            text = para.text.strip()
            if text:
                fmt["headings"].append({"level": level, "text": text[:50]})

    return fmt


def extract_text(doc):
    """提取文档文本"""
    paragraphs = []
    for para in doc.paragraphs:
        text = para.text.strip()
        if text:
            paragraphs.append(text)
    return {
        "paragraph_count": len(paragraphs),
        "total_chars": sum(len(p) for p in paragraphs),
        "paragraphs": paragraphs,
        "first_sentences": [p[:50] for p in paragraphs[:10] if p],
        "last_sentences": [p[-50:] for p in paragraphs if p][-5:],
    }


def process_file(file_path):
    """处理单个 docx 文件"""
    try:
        doc = Document(str(file_path))
        return {
            "file": os.path.basename(file_path),
            "file_type": "docx",
            "format_profile": extract_format(doc),
            "text_profile": extract_text(doc),
            "status": "success",
        }
    except Exception as e:
        return {
            "file": os.path.basename(file_path),
            "file_type": "docx",
            "status": "failed",
            "error": str(e),
        }


def main():
    parser = argparse.ArgumentParser(description="从 docx 提取格式与文本")
    parser.add_argument("--input", required=True, help="输入文件或目录")
    parser.add_argument("--output", required=True, help="输出 JSON 文件")
    args = parser.parse_args()

    input_path = Path(args.input)
    if input_path.is_dir():
        files = list(input_path.glob("**/*.docx"))
    else:
        files = [input_path]

    results = [process_file(f) for f in files]

    with open(args.output, "w", encoding="utf-8") as f:
        json.dump(results, f, ensure_ascii=False, indent=2)

    success_count = sum(1 for r in results if r["status"] == "success")
    print(f"✅ 处理完成：{success_count}/{len(results)} 成功，输出至 {args.output}")


if __name__ == "__main__":
    main()