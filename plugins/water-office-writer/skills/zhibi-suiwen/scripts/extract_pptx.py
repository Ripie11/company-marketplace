#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
extract_pptx.py - 从 .pptx 提取文本与版式

用法:
    python extract_pptx.py --input <文件或目录> --output <JSON输出文件>

依赖:
    pip install python-pptx
"""
import argparse
import json
import os
from pathlib import Path
from pptx import Presentation
from pptx.util import Emu


def emu_to_pt(emu):
    return round(emu / 12700, 2) if emu else None


def extract_pptx(file_path):
    """提取单个 pptx 文件"""
    try:
        prs = Presentation(str(file_path))
        result = {
            "file": os.path.basename(file_path),
            "file_type": "pptx",
            "slide_count": len(prs.slides),
            "slide_width_emu": prs.slide_width,
            "slide_height_emu": prs.slide_height,
            "slides": [],
        }

        all_text = []
        for i, slide in enumerate(prs.slides):
            slide_data = {"index": i + 1, "shapes": [], "notes": ""}
            for shape in slide.shapes:
                if shape.has_text_frame:
                    for para in shape.text_frame.paragraphs:
                        text = "".join(run.text for run in para.runs).strip()
                        if text:
                            slide_data["shapes"].append({
                                "text": text,
                                "font_name": para.runs[0].font.name if para.runs else None,
                                "font_size_pt": para.runs[0].font.size.pt if para.runs and para.runs[0].font.size else None,
                            })
                            all_text.append(text)
            if slide.has_notes_slide:
                notes_text = slide.notes_slide.notes_text_frame.text
                if notes_text:
                    slide_data["notes"] = notes_text
                    all_text.append(notes_text)
            result["slides"].append(slide_data)

        result["text_profile"] = {
            "total_chars": sum(len(t) for t in all_text),
            "paragraph_count": len(all_text),
            "first_sentences": [t[:50] for t in all_text[:10] if t],
            "all_text_sample": all_text[:30],  # 仅保留前30条用于风格分析
        }
        result["status"] = "success"
        return result
    except Exception as e:
        return {
            "file": os.path.basename(file_path),
            "file_type": "pptx",
            "status": "failed",
            "error": str(e),
        }


def main():
    parser = argparse.ArgumentParser(description="从 pptx 提取文本与版式")
    parser.add_argument("--input", required=True, help="输入文件或目录")
    parser.add_argument("--output", required=True, help="输出 JSON 文件")
    args = parser.parse_args()

    input_path = Path(args.input)
    if input_path.is_dir():
        files = list(input_path.glob("**/*.pptx"))
    else:
        files = [input_path]

    # 跳过 Office 临时文件
    files = [f for f in files if not f.name.startswith("~$")]

    results = [extract_pptx(f) for f in files]

    with open(args.output, "w", encoding="utf-8") as f:
        json.dump(results, f, ensure_ascii=False, indent=2)

    success_count = sum(1 for r in results if r["status"] == "success")
    print(f"✅ 处理完成：{success_count}/{len(results)} 成功，输出至 {args.output}")


if __name__ == "__main__":
    main()