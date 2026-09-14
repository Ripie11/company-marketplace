#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""audit_outputs.py — 交付前产物审计：纸张/表格/页码/标记残留"""
import glob
import os
import sys
from docx import Document

D = sys.argv[1] if len(sys.argv) > 1 else "out_final"

MARKERS = [("管道符|", "|"), ("加粗**", "**"),
           ("井号#", "#"), ("列表'- '", "\n- ")]

rows = []
tot_tbl = tot_cells = 0
for f in sorted(glob.glob(os.path.join(D, "C*.docx"))):
    d = Document(f)
    s = d.sections[0]
    txt = "\n".join(p.text for p in d.paragraphs)
    haspg = any(b"PAGE" in getattr(pt, "blob", b"")
                for pt in d.part.package.iter_parts())
    cells = sum(len(r.cells) for t in d.tables for r in t.rows)
    tot_tbl += len(d.tables)
    tot_cells += cells
    resid = [name for name, m in MARKERS if m in txt]
    rows.append((os.path.basename(f)[:-5][:24],
                 f"{s.page_width.cm:.1f}x{s.page_height.cm:.1f}",
                 len(d.tables), cells, haspg,
                 "、".join(resid) if resid else "无"))

hdr = f'{"用例":26s}{"纸张(cm)":>10s}{"表":>4s}{"格":>5s}{"PAGE":>6s}  标记残留'
print(hdr)
print("-" * len(hdr))
for r in rows:
    print(f"{r[0]:26s}{r[1]:>10s}{r[2]:4d}{r[3]:5d}{('有' if r[4] else '无'):>7s}  {r[5]}")
print(f"\n合计：{len(rows)} 份 docx | 表格 {tot_tbl} 个 / {tot_cells} 格 | "
      f"全部 A4={all(r[1].startswith('21.0x29.7') for r in rows)} | "
      f"全部有页码={all(r[4] for r in rows)} | "
      f"标记零残留={all(r[5] == '无' for r in rows)}")
